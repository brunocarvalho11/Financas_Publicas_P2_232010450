"""Reproduz as tabelas do artigo (2014-2023) a partir dos arquivos do repositório.

Uso (na raiz do repositório):  python 4_Codigo/analise.py
Saídas: 3_Dados_Tratados/series_reais_2014_2023.csv e 3_Dados_Tratados/resultados.txt
"""
import csv
from pathlib import Path

import numpy as np
import openpyxl
from scipy import stats

RAIZ = Path(__file__).resolve().parent.parent
BASE_P1 = RAIZ / "1_Dados_Originais_P1" / "Dados_Financas_P1_241021847.xlsx"
PIB_XLSX = RAIZ / "2_Dados_Brutos_Atualizados" / "PIB_MUNICIPIOS.xlsx"
IPCA_XLSX = RAIZ / "2_Dados_Brutos_Atualizados" / "IPCA.xlsx"
SAIDA = RAIZ / "3_Dados_Tratados"
ANOS = list(range(2014, 2024))
ANO_BASE = 2023


def ipca_dezembro():
    ws = openpyxl.load_workbook(IPCA_XLSX, data_only=True)["Séries"]
    d = {}
    for data, valor, *_ in ws.iter_rows(min_row=2, values_only=True):
        if data and str(data).endswith(".12") and valor is not None:
            d[int(str(data).split(".")[0])] = float(valor)
    return d


def pib_municipal(nome):
    ws = openpyxl.load_workbook(PIB_XLSX, data_only=True)["Tabela"]
    for r in ws.iter_rows(values_only=True):
        if r[0] and r[0].strip() == nome:
            return [float(x) for x in r[1:11]]  # mil R$ correntes, 2014-2023
    raise ValueError(nome)


def aba(wb, nome):
    return [r for r in wb[nome].iter_rows(min_row=2, max_row=11, values_only=True)]


def reg(x, y):
    s = stats.linregress(x, y)
    return s.slope, s.stderr, s.pvalue, s.rvalue ** 2


def main():
    wb = openpyxl.load_workbook(BASE_P1, data_only=True)
    ipca = ipca_dezembro()
    fator = np.array([ipca[ANO_BASE] / ipca[a] for a in ANOS])
    out, linhas = [], []

    entes = {
        "Betim": ("Betim-MG (Micro 1)", "Betim"),
        "Governador Valadares": ("Gov. Valadares-MG (Micro 2)", "Governador Valadares"),
    }
    series = {}
    for ente, (nome_aba, nome_pib) in entes.items():
        d = aba(wb, nome_aba)
        pop = np.array([r[1] for r in d], float)
        Rn = np.array([r[3] for r in d], float)
        Dn = np.array([r[4] for r in d], float)
        R, D = Rn * fator, Dn * fator
        pib_n = np.array(pib_municipal(nome_pib))
        pib_r = pib_n * 1000 * fator
        series[ente] = dict(R=R, D=D, pop=pop, pib=pib_r, pib_n=pib_n * 1000)
        for i, a in enumerate(ANOS):
            out.append([ente, a, pop[i], Rn[i], Dn[i], round(fator[i], 6), R[i], D[i], pib_n[i] * 1000, pib_r[i]])

    d = aba(wb, "Indonésia (Macro)")
    R = np.array([r[7] for r in d], float)
    D = np.array([r[8] for r in d], float)
    series["Indonésia"] = dict(R=R, D=D)
    for i, a in enumerate(ANOS):
        out.append(["Indonésia", a, "", d[i][3], d[i][4], d[i][5], R[i], D[i], "", ""])

    SAIDA.mkdir(exist_ok=True)
    with open(SAIDA / "series_reais_2014_2023.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["ente", "ano", "populacao", "receita_nominal", "despesa_nominal", "fator_deflacao",
                    "receita_real_2023", "despesa_real_2023", "pib_nominal_rs", "pib_real_2023_rs"])
        w.writerows(out)

    p = linhas.append
    for ente, s in series.items():
        R, D = s["R"], s["D"]
        p(f"===== {ente} =====")
        p(f"Receita real: {R[0]:.3e} -> {R[-1]:.3e} ({100*(R[-1]/R[0]-1):.1f}%; {100*((R[-1]/R[0])**(1/9)-1):.2f}% a.a.)")
        p(f"Despesa real: {D[0]:.3e} -> {D[-1]:.3e} ({100*(D[-1]/D[0]-1):.1f}%; {100*((D[-1]/D[0])**(1/9)-1):.2f}% a.a.)")
        p("Saldo/Receita (%): " + ", ".join(f"{x:.1f}" for x in 100 * (R - D) / R) + f" | anos de deficit: {(R < D).sum()}")
        r, pv = stats.pearsonr(R, D)
        p(f"Correlacao em nivel R,D: {r:.3f} (p={pv:.4f})")
        dR, dD = np.diff(np.log(R)), np.diff(np.log(D))
        r, pv = stats.pearsonr(dR, dD)
        b, se, pb, r2 = reg(dR, dD)
        p(f"Correlacao em variacoes: {r:.3f} (p={pv:.3f}); dlnD = a + b dlnR: b={b:.3f} (ep {se:.3f}; p={pb:.3f})")
        p(f"Desvio-padrao dlnR={dR.std(ddof=1):.3f}  dlnD={dD.std(ddof=1):.3f}")
        if "pib" in s:
            P, pop = s["pib"], s["pop"]
            p("Receita pc real: " + ", ".join(f"{x:,.0f}" for x in R / pop))
            p("Despesa pc real: " + ", ".join(f"{x:,.0f}" for x in D / pop))
            p("Receita/PIB (%): " + ", ".join(f"{x:.2f}" for x in 100 * R / P))
            b, se, pb, r2 = reg(P, R)
            p(f"NIVEL  R = a + b PIB(real): b={b:.4f}  R2={r2:.3f}  p={pb:.4f}")
            b, se, pb, r2 = reg(np.log(P), np.log(R))
            p(f"LOGLOG elasticidade={b:.3f} (ep {se:.3f}; p={pb:.4f}) R2={r2:.3f}")
            b, se, pb, r2 = reg(np.diff(np.log(P)), dR)
            p(f"DIFERENCAS dlnR = a + b dlnPIB: b={b:.3f} (ep {se:.3f}; p={pb:.3f})")
            # Conferência do resultado da P1 (PIB nominal em mil R$ x receita real)
            b, se, pb, r2 = reg(s["pib_n"] / 1000, R)
            p(f"[CONFERENCIA] receita REAL x PIB NOMINAL (mil R$): b={b:.2f}  R2={r2:.4f}  "
              "(combinacao inconsistente; a P1 usou-a para Valadares: b=133,90 e R2=0,8025)")
        p("")
    (SAIDA / "resultados.txt").write_text("\n".join(linhas), encoding="utf-8")
    print("\n".join(linhas))


if __name__ == "__main__":
    main()
